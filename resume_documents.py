"""Privacy-first resume document extraction.

The Telegram flow may accept a PDF or DOCX for Resume Match, but the bytes are
never written to disk by this module.  Parsing is intentionally bounded to keep
an untrusted upload from becoming an availability problem.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_DOCX_MEMBERS = 1000
MAX_PDF_PAGES = 25
MAX_TEXT_CHARS = 30_000
MIN_TEXT_CHARS = 120


class ResumeDocumentError(ValueError):
    """Safe, user-displayable failure raised for unsupported/bad uploads."""


@dataclass(frozen=True)
class ResumeDocumentText:
    kind: str
    text: str


def _normalize_text(text: str) -> str:
    text = str(text or "").replace("\x00", " ")
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if line:
            lines.append(line)
    normalized = "\n".join(lines).strip()
    return normalized[:MAX_TEXT_CHARS]


def _looks_like_pdf(data: bytes) -> bool:
    # PDF headers should be near the beginning; tolerate a small binary prefix.
    return b"%PDF-" in data[:1024]


def _validate_docx_container(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_MEMBERS:
                raise ResumeDocumentError("DOCX слишком сложный для безопасной обработки.")
            uncompressed = sum(max(0, int(info.file_size)) for info in infos)
            if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeDocumentError("DOCX слишком большой после распаковки.")
            names = {info.filename for info in infos}
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(names):
                raise ResumeDocumentError("Файл не похож на корректный DOCX.")
    except ResumeDocumentError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ResumeDocumentError("Файл не похож на корректный DOCX.") from exc


def _detect_kind(data: bytes, filename: str, mime_type: str) -> str:
    suffix = Path(str(filename or "")).suffix.lower()
    mime = str(mime_type or "").lower().strip()

    pdf_hint = suffix == ".pdf" or mime == PDF_MIME
    docx_hint = suffix == ".docx" or mime == DOCX_MIME

    if pdf_hint and _looks_like_pdf(data):
        return "pdf"
    if docx_hint and data.startswith(b"PK"):
        _validate_docx_container(data)
        return "docx"

    # MIME types are user-controlled in Telegram.  Do not trust an extension or
    # mime claim when the actual container signature disagrees.
    if pdf_hint:
        raise ResumeDocumentError("Файл подписан как PDF, но его содержимое не похоже на PDF.")
    if docx_hint:
        raise ResumeDocumentError("Файл подписан как DOCX, но его содержимое не похоже на DOCX.")
    raise ResumeDocumentError("Поддерживаются только PDF и DOCX.")


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
    except Exception as exc:
        raise ResumeDocumentError("Не удалось прочитать PDF.") from exc

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            raise ResumeDocumentError("PDF защищён паролем. Пришли незашифрованный файл или текст.")

    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        raise ResumeDocumentError(f"PDF слишком длинный: максимум {MAX_PDF_PAGES} страниц.")

    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # One malformed page should not discard readable text from the rest.
            continue
    text = _normalize_text("\n".join(parts))
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeDocumentError(
            "В PDF почти нет извлекаемого текста. Если это скан, пришли DOCX или вставь текст резюме."
        )
    return text


def _extract_docx(data: bytes) -> str:
    _validate_docx_container(data)
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise ResumeDocumentError("Не удалось прочитать DOCX.") from exc

    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
    # CV templates frequently put the useful content into tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    text = _normalize_text("\n".join(parts))
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeDocumentError("В DOCX слишком мало текста для сравнения с вакансией.")
    return text


def extract_resume_document(data: bytes | bytearray, filename: str = "", mime_type: str = "") -> ResumeDocumentText:
    """Extract bounded plain text from an in-memory PDF/DOCX upload.

    The caller is expected to discard the returned text after one Resume Match.
    No filename, bytes or extracted text are logged or persisted here.
    """
    payload = bytes(data or b"")
    if not payload:
        raise ResumeDocumentError("Файл пустой.")
    if len(payload) > MAX_FILE_BYTES:
        raise ResumeDocumentError("Файл слишком большой: максимум 5 МБ.")

    kind = _detect_kind(payload, filename, mime_type)
    if kind == "pdf":
        text = _extract_pdf(payload)
    else:
        text = _extract_docx(payload)
    return ResumeDocumentText(kind=kind, text=text)
